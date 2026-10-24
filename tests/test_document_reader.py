import io
import unittest

from docx import Document
from pypdf import PdfWriter

from document_reader import MAX_FILE_SIZE_BYTES, MAX_TEXT_CHARACTERS, extract_text


class UploadedFileStub:
    def __init__(self, name, content):
        self.name = name
        self._content = content

    def getvalue(self):
        return self._content


class DocumentReaderTests(unittest.TestCase):
    def test_extracts_txt_content(self):
        uploaded_file = UploadedFileStub("notes.txt", b"Important project notes")

        result = extract_text(uploaded_file)

        self.assertEqual(result["text"], "Important project notes")
        self.assertEqual(result["file_type"], "TXT")

    def test_extracts_docx_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph("Project requirements")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Owner"
        table.cell(0, 1).text = "Velugu Rishik Gani"
        buffer = io.BytesIO()
        document.save(buffer)

        result = extract_text(UploadedFileStub("requirements.docx", buffer.getvalue()))

        self.assertIn("Project requirements", result["text"])
        self.assertIn("Owner | Velugu Rishik Gani", result["text"])

    def test_rejects_empty_files(self):
        with self.assertRaisesRegex(ValueError, "empty"):
            extract_text(UploadedFileStub("empty.txt", b""))

    def test_rejects_unsupported_formats(self):
        with self.assertRaisesRegex(ValueError, "Unsupported file type"):
            extract_text(UploadedFileStub("image.png", b"not a document"))

    def test_rejects_files_over_size_limit(self):
        oversized_content = b"x" * (MAX_FILE_SIZE_BYTES + 1)

        with self.assertRaisesRegex(ValueError, "too large"):
            extract_text(UploadedFileStub("large.txt", oversized_content))

    def test_rejects_extracted_text_over_character_limit(self):
        oversized_text = ("x" * MAX_TEXT_CHARACTERS) + "y"

        with self.assertRaisesRegex(ValueError, "too much text"):
            extract_text(UploadedFileStub("large.txt", oversized_text.encode()))

    def test_rejects_pdf_without_readable_text(self):
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        buffer = io.BytesIO()
        writer.write(buffer)

        with self.assertRaisesRegex(ValueError, "No readable text"):
            extract_text(UploadedFileStub("scanned.pdf", buffer.getvalue()))


if __name__ == "__main__":
    unittest.main()
