from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)


def create_report(comparison_result, document_a_name, document_b_name):
    """Create a Word report from the comparison result."""

    report = Document()

    report.add_heading(
        "AI Document Comparison Report",
        level=0,
    )

    report.add_paragraph(
        f"Document A: {document_a_name}"
    )

    report.add_paragraph(
        f"Document B: {document_b_name}"
    )

    report.add_heading("Document A summary", level=1)
    report.add_paragraph(
        comparison_result["document_a_summary"]
    )

    report.add_heading("Document B summary", level=1)
    report.add_paragraph(
        comparison_result["document_b_summary"]
    )

    report.add_heading("Similarities", level=1)

    for similarity in comparison_result["similarities"]:
        report.add_paragraph(
            similarity,
            style="List Bullet",
        )

    report.add_heading("Differences", level=1)

    for difference in comparison_result["differences"]:
        if isinstance(difference, dict):
            title = difference.get("title", "Document change")
            category = difference.get("category", "Change")
            severity = difference.get("severity", "Medium")
            description = difference.get(
                "description",
                "No explanation provided.",
            )
            document_a = difference.get(
                "document_a",
                "Not found in the provided text.",
            )
            document_b = difference.get(
                "document_b",
                "Not found in the provided text.",
            )

            report.add_paragraph(
                f"{title} ({category}; {severity})",
                style="List Bullet",
            )
            report.add_paragraph(description)
            report.add_paragraph(f"Document A: {document_a}")
            report.add_paragraph(f"Document B: {document_b}")
        else:
            report.add_paragraph(
                str(difference),
                style="List Bullet",
            )

    report.add_heading("Key details", level=1)

    for detail in comparison_result["key_details"]:
        report.add_paragraph(
            f"Field: {detail.get('field', 'Unknown')}",
            style="List Bullet",
        )
        report.add_paragraph(
            f"Document A: {detail.get('document_a', 'Not found')}"
        )
        report.add_paragraph(
            f"Document B: {detail.get('document_b', 'Not found')}"
        )

    report.add_heading("Overall conclusion", level=1)
    report.add_paragraph(
        comparison_result["overall_conclusion"]
    )

    report.add_heading("Limitations", level=1)

    for limitation in comparison_result["limitations"]:
        report.add_paragraph(
            limitation,
            style="List Bullet",
        )

    report.add_heading("Risk flags", level=1)
    for risk in comparison_result.get("risk_flags", []):
        report.add_paragraph(risk, style="List Bullet")

    report.add_heading("Deadlines", level=1)
    for deadline in comparison_result.get("deadlines", []):
        report.add_paragraph(deadline, style="List Bullet")

    report.add_heading("Source citations", level=1)
    for citation in comparison_result.get("source_citations", []):
        report.add_paragraph(
            f"{citation.get('finding', 'Finding')} ({citation.get('source', 'Source')})",
            style="List Bullet",
        )
        report.add_paragraph(f"\"{citation.get('quote', 'Quote unavailable')}\"")

    output = BytesIO()
    report.save(output)
    output.seek(0)

    return output


def create_pdf_report(comparison_result, document_a_name, document_b_name):
    """Create a PDF report from the comparison result."""

    output = BytesIO()
    report = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
        title="AI Document Comparison Report",
    )

    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ComparisonTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            textColor=colors.HexColor("#13324F"),
            fontSize=18,
            leading=22,
            spaceAfter=16,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ComparisonHeading",
            parent=styles["Heading2"],
            textColor=colors.HexColor("#13324F"),
            spaceBefore=12,
            spaceAfter=6,
            keepWithNext=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ComparisonBody",
            parent=styles["BodyText"],
            fontSize=9,
            leading=13,
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ComparisonMeta",
            parent=styles["BodyText"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#394B59"),
            spaceAfter=4,
        )
    )

    def safe_text(value):
        if value is None:
            return ""
        return escape(str(value)).replace("\n", "<br/>")

    def add_paragraph(text, style="ComparisonBody"):
        story.append(Paragraph(safe_text(text), styles[style]))

    def add_section(title, items):
        add_paragraph(title, "ComparisonHeading")
        if not items:
            add_paragraph("None identified.", "ComparisonMeta")
            return
        for item in items:
            story.append(
                Paragraph(f"&bull; {safe_text(item)}", styles["ComparisonBody"])
            )

    story = [
        Paragraph("AI Document Comparison Report", styles["ComparisonTitle"]),
        Paragraph(f"<b>Document A:</b> {safe_text(document_a_name)}", styles["ComparisonMeta"]),
        Paragraph(f"<b>Document B:</b> {safe_text(document_b_name)}", styles["ComparisonMeta"]),
        Spacer(1, 8),
    ]

    add_paragraph("Document A summary", "ComparisonHeading")
    add_paragraph(comparison_result.get("document_a_summary", ""))
    add_paragraph("Document B summary", "ComparisonHeading")
    add_paragraph(comparison_result.get("document_b_summary", ""))

    add_section("Similarities", comparison_result.get("similarities", []))

    differences = comparison_result.get("differences", [])
    add_paragraph("Differences", "ComparisonHeading")
    if differences:
        for difference in differences:
            if isinstance(difference, dict):
                title = difference.get("title", "Document change")
                category = difference.get("category", "Change")
                severity = difference.get("severity", "Medium")
                add_paragraph(f"{title} ({category}; {severity})")
                add_paragraph(difference.get("description", "No explanation provided."))
                add_paragraph(
                    f"Document A: {difference.get('document_a', 'Not found in the provided text.')}",
                    "ComparisonMeta",
                )
                add_paragraph(
                    f"Document B: {difference.get('document_b', 'Not found in the provided text.')}",
                    "ComparisonMeta",
                )
            else:
                story.append(
                    Paragraph(
                        f"&bull; {safe_text(difference)}",
                        styles["ComparisonBody"],
                    )
                )
    else:
        add_paragraph("None identified.", "ComparisonMeta")

    key_details = comparison_result.get("key_details", [])
    add_paragraph("Key details", "ComparisonHeading")
    if key_details:
        for detail in key_details:
            add_paragraph(f"Field: {detail.get('field', 'Unknown')}")
            add_paragraph(
                f"Document A: {detail.get('document_a', 'Not found')}",
                "ComparisonMeta",
            )
            add_paragraph(
                f"Document B: {detail.get('document_b', 'Not found')}",
                "ComparisonMeta",
            )
    else:
        add_paragraph("None identified.", "ComparisonMeta")

    add_paragraph("Overall conclusion", "ComparisonHeading")
    add_paragraph(comparison_result.get("overall_conclusion", ""))
    add_section("Limitations", comparison_result.get("limitations", []))
    add_section("Risk flags", comparison_result.get("risk_flags", []))
    add_section("Deadlines", comparison_result.get("deadlines", []))

    citations = comparison_result.get("source_citations", [])
    add_paragraph("Source citations", "ComparisonHeading")
    if citations:
        for citation in citations:
            add_paragraph(
                f"{citation.get('finding', 'Finding')} "
                f"({citation.get('source', 'Source')})"
            )
            add_paragraph(
                f"“{citation.get('quote', 'Quote unavailable')}”",
                "ComparisonMeta",
            )
    else:
        add_paragraph("None provided.", "ComparisonMeta")

    def add_page_number(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawRightString(
            letter[0] - 0.75 * inch,
            0.4 * inch,
            f"Page {document.page}",
        )
        canvas.restoreState()

    report.build(
        story,
        onFirstPage=add_page_number,
        onLaterPages=add_page_number,
    )
    output.seek(0)
    return output
