from io import BytesIO

from docx import Document
from pypdf import PdfReader


MAX_FILE_SIZE_BYTES = 25 * 1024 * 1024
MAX_TEXT_CHARACTERS = 120_000


def _validate_text_length(text, file_name):
    if len(text) > MAX_TEXT_CHARACTERS:
        raise ValueError(
            f"{file_name} contains too much text ({len(text):,} characters). "
            f"Please upload a shorter document or split it into parts. "
            f"The limit is {MAX_TEXT_CHARACTERS:,} characters."
        )


def extract_text(uploaded_file):
    """Extract readable text from a PDF, DOCX, or TXT file."""

    if uploaded_file is None:
        raise ValueError("Please upload a document.")

    file_name = uploaded_file.name
    extension = file_name.lower().rsplit(".", 1)[-1]
    file_bytes = uploaded_file.getvalue()

    if not file_bytes:
        raise ValueError(f"The uploaded file is empty: {file_name}")

    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        size_mb = len(file_bytes) / (1024 * 1024)
        raise ValueError(
            f"{file_name} is too large ({size_mb:.1f} MB). "
            "Please upload a file smaller than 25 MB."
        )

    if extension == "pdf":
        return _extract_pdf(file_bytes)

    if extension == "docx":
        return _extract_docx(file_bytes)

    if extension == "txt":
        return _extract_txt(file_bytes)

    raise ValueError(
        "Unsupported file type. Please use PDF, DOCX, or TXT files."
    )


def _extract_pdf(file_bytes):
    try:
        pdf_reader = PdfReader(BytesIO(file_bytes))
    except Exception as error:
        raise ValueError("This PDF could not be opened.") from error

    page_texts = []

    for page_number, page in enumerate(pdf_reader.pages, start=1):
        text = page.extract_text() or ""

        if text.strip():
            page_texts.append(
                f"[Page {page_number}]\n{text.strip()}"
            )

    full_text = "\n\n".join(page_texts)

    if not full_text.strip():
        raise ValueError(
            "No readable text was found in this PDF. "
            "It may be a scanned PDF."
        )

    _validate_text_length(full_text, "The extracted PDF text")

    return {
        "text": full_text,
        "unit_label": "Pages",
        "unit_count": len(pdf_reader.pages),
        "character_count": len(full_text),
        "file_type": "PDF",
    }


def _extract_docx(file_bytes):
    try:
        document = Document(BytesIO(file_bytes))
    except Exception as error:
        raise ValueError("This DOCX file could not be opened.") from error

    parts = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)

            if row_text:
                parts.append(row_text)

    full_text = "\n\n".join(parts)

    if not full_text.strip():
        raise ValueError("No readable text was found in this DOCX file.")

    _validate_text_length(full_text, "The extracted DOCX text")

    return {
        "text": full_text,
        "unit_label": "Paragraphs and table rows",
        "unit_count": len(parts),
        "character_count": len(full_text),
        "file_type": "DOCX",
    }


def _extract_txt(file_bytes):
    full_text = file_bytes.decode("utf-8", errors="replace").strip()

    if not full_text:
        raise ValueError("No readable text was found in this TXT file.")

    _validate_text_length(full_text, "The TXT file")

    return {
        "text": full_text,
        "unit_label": "Lines",
        "unit_count": len(full_text.splitlines()),
        "character_count": len(full_text),
        "file_type": "TXT",
    }
